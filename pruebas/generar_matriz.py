"""
Ejecuta las pruebas funcionales automatizadas y genera la matriz de resultados.

    python pruebas/generar_matriz.py [--responsable "Nombre"]

Salidas:
    docs/pruebas/matriz_pruebas_funcionales.xlsx
    docs/pruebas/matriz_pruebas_funcionales.md

Cada fila proviene del docstring de una prueba («Requisito | Módulo |
Precondiciones | Pasos | Resultado esperado») y del resultado real obtenido.
"""
import argparse
import os
import sys
import unittest
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "sgr.settings")
os.environ.setdefault("DJANGO_DEBUG", "1")

import django  # noqa: E402

django.setup()

from django.db import connection  # noqa: E402
from django.test.runner import DiscoverRunner  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

COLUMNAS = ["ID", "Requisito", "Módulo", "Precondiciones", "Pasos", "Resultado esperado", "Resultado obtenido",
            "Estado", "Evidencia", "Responsable", "Fecha"]


class Resultado(unittest.TextTestResult):
    registros = []

    def _registrar(self, test, estado, detalle=""):
        self.registros.append((test, estado, detalle))

    def addSuccess(self, test):
        super().addSuccess(test)
        self._registrar(test, "Aprobada")

    def addFailure(self, test, err):
        super().addFailure(test, err)
        self._registrar(test, "Fallida", self._exc_info_to_string(err, test).strip().splitlines()[-1])

    def addError(self, test, err):
        super().addError(test, err)
        self._registrar(test, "Error", self._exc_info_to_string(err, test).strip().splitlines()[-1])

    def addSkip(self, test, reason):
        super().addSkip(test, reason)
        self._registrar(test, "Pendiente", reason)


class Runner(DiscoverRunner):
    def get_resultclass(self):
        return Resultado


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--responsable", default=os.environ.get("SGR_RESPONSABLE_PRUEBAS", "Equipo SGR"))
    args = parser.parse_args()

    Resultado.registros = []
    runner = Runner(verbosity=1)
    fallas = runner.run_tests(["tests"])
    fecha = datetime.now().strftime("%d-%m-%Y %H:%M")
    motor = connection.vendor
    filas = []
    for i, (test, estado, detalle) in enumerate(sorted(Resultado.registros, key=lambda r: r[0].id()), start=1):
        doc = (test._testMethodDoc or "").strip()
        partes = [p.strip() for p in doc.split("|")]
        partes += [""] * (5 - len(partes))
        requisito, modulo, pre, pasos, esperado = partes[:5]
        obtenido = "Coincide con lo esperado" if estado == "Aprobada" else detalle
        filas.append([f"PF-{i:03d}", requisito, modulo, pre, pasos, esperado, obtenido, estado,
                      f"Prueba automatizada: {test.id()}", args.responsable, fecha])

    destino = BASE / "docs" / "pruebas"
    destino.mkdir(parents=True, exist_ok=True)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pruebas funcionales"
    hoja.append(COLUMNAS)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="C41230")
    for f in filas:
        hoja.append(f)
    anchos = [9, 26, 20, 34, 44, 44, 26, 11, 60, 16, 17]
    for i, ancho in enumerate(anchos):
        hoja.column_dimensions[chr(65 + i)].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    resumen = libro.create_sheet("Resumen")
    aprobadas = sum(1 for f in filas if f[7] == "Aprobada")
    for fila in [["Fecha de ejecución", fecha], ["Total de pruebas", len(filas)], ["Aprobadas", aprobadas],
                 ["No aprobadas", len(filas) - aprobadas], ["Motor de base de datos", motor], ["Comando", "python pruebas/generar_matriz.py"]]:
        resumen.append(fila)
    libro.save(destino / "matriz_pruebas_funcionales.xlsx")

    md = ["# Matriz de pruebas funcionales", "",
          f"Ejecución: **{fecha}** · Total: **{len(filas)}** · Aprobadas: **{aprobadas}** · "
          f"No aprobadas: **{len(filas) - aprobadas}** · Base de datos: **{motor}**", "",
          "Generado automáticamente con `python pruebas/generar_matriz.py` (no editar a mano).", "",
          "| ID | Requisito | Módulo | Precondiciones | Pasos | Resultado esperado | Estado | Evidencia |",
          "|---|---|---|---|---|---|---|---|"]
    for f in filas:
        celdas = [f[0], f[1], f[2], f[3], f[4], f[5], f[7], "`" + f[8].replace("Prueba automatizada: ", "") + "`"]
        md.append("| " + " | ".join(c.replace("|", "/") for c in celdas) + " |")
    (destino / "matriz_pruebas_funcionales.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print(f"Matriz generada en {destino} ({aprobadas}/{len(filas)} aprobadas).")
    sys.exit(bool(fallas))


if __name__ == "__main__":
    main()
